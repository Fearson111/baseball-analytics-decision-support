from __future__ import annotations

import html
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys

import pandas as pd

from ..table_display import display_value as table_value, decimal_places, display_name
from ..paths import configured_path
from .markets import key, number


PAPER = "FFC3A0"
TEAMS = {
    "ARI": ("A71930", "FFFFFF"), "AZ": ("A71930", "FFFFFF"), "ATL": ("CE1141", "FFFFFF"),
    "BAL": ("DF4601", "000000"), "BOS": ("BD3039", "FFFFFF"), "CHC": ("0E3386", "FFFFFF"),
    "CWS": ("27251F", "FFFFFF"), "CHW": ("27251F", "FFFFFF"), "CIN": ("C6011F", "FFFFFF"),
    "CLE": ("00385D", "FFFFFF"), "COL": ("333366", "FFFFFF"), "DET": ("0C2340", "FFFFFF"),
    "HOU": ("EB6E1F", "000000"), "KC": ("004687", "FFFFFF"), "KCR": ("004687", "FFFFFF"),
    "LAA": ("BA0021", "FFFFFF"), "LAD": ("005A9C", "FFFFFF"), "MIA": ("00A3E0", "000000"),
    "MIL": ("12284B", "FFFFFF"), "MIN": ("002B5C", "FFFFFF"), "NYM": ("002D72", "FFFFFF"),
    "NYY": ("0C2340", "FFFFFF"), "ATH": ("003831", "FFFFFF"), "OAK": ("003831", "FFFFFF"),
    "PHI": ("E81828", "FFFFFF"), "PIT": ("FDB827", "000000"), "SD": ("FFC425", "000000"),
    "SDP": ("FFC425", "000000"), "SEA": ("005C5C", "FFFFFF"), "SF": ("FD5A1E", "000000"),
    "SFG": ("FD5A1E", "000000"), "STL": ("C41E3A", "FFFFFF"), "TB": ("092C5C", "FFFFFF"),
    "TBR": ("092C5C", "FFFFFF"), "TEX": ("003278", "FFFFFF"), "TOR": ("134A8E", "FFFFFF"),
    "WSH": ("AB0003", "FFFFFF"), "WSN": ("AB0003", "FFFFFF"),
}

VIEWS = {
    "All-Seeing Eye": [("eye_rank", "rank"), ("player", "player"), ("eye_score", "eye_score"),
                      ("gbm_rank", "model_1"), ("gbm_progression_rank", "regression"),
                      ("sequence_rank", "model_2"), ("pitcher_batter_rank", "matchup_rank")],
    "GBM Top Hitters": [("rank", "rk"), ("player_name", "player"), ("team_abbrev", "tm"),
                        ("adjusted_hr_probability_percent", "model_hr%"), ("american_betting_odds", "model_odds"),
                        ("book_american_odds", "book_odds"), ("edge_percent", "edge%")],
    "Positive Regression": [("regression_rank", "rk"), ("player_name", "player"),
                            ("positive_regression_score", "regr_score"), ("american_betting_odds", "model_odds"),
                            ("book_american_odds", "book_odds"), ("edge_percent", "edge%")],
    "Matchups": [("rank", "rank"), ("player_name", "player_name"), ("opposing_pitcher", "opposing_pitcher"),
                 ("score", "score"), ("hr_edge", "hr_edge")],
    "Straight Bets": [("signal", "signal"), ("player", "player"), ("odds", "odds"), ("signal_stats", "signal_history")],
}


def posting_frame(frame: pd.DataFrame, view: str, limit: int = 25) -> pd.DataFrame:
    source = frame.copy().fillna("")
    if view == "Straight Bets":
        source = source.sort_values(["player", "signal"], kind="stable")
    elif limit:
        source = source.head(limit)
    spec = VIEWS.get(view, [(col, col) for col in source.columns if col != "team"])
    output = pd.DataFrame(index=source.index)
    for column, label in spec:
        output[label] = source.get(column, pd.Series("", index=source.index))
    output = output.map(lambda value: display_name(value) if isinstance(value, str) else value)
    output = output.reset_index(drop=True)
    output.attrs["teams"] = source.get("team", source.get("team_abbrev", pd.Series("", index=source.index))).tolist()
    return output


def display_value(value: object, column: str) -> str:
    if value is None or str(value).strip().lower() in {"nan", "none", "<na>", ""}:
        return ""
    num = number(value)
    if num is None:
        return display_name(str(value))
    if column in {"book_odds", "model_odds", "parlay_odds"}:
        return f"{num:+.0f}"
    return table_value(value, column)


def grouped_name(frame: pd.DataFrame, row: int, col: str) -> bool:
    return col == "player" and "signal" in frame and row > 0 and frame.iloc[row][col] == frame.iloc[row - 1][col]


def table_html(frame: pd.DataFrame) -> str:
    teams = frame.attrs.get("teams", [])
    parts = ['<div class="posting-table">',
             '<table style="border-collapse:collapse;font-family:Menlo,monospace;font-size:11pt;letter-spacing:0;white-space:nowrap">',
             '<thead><tr>']
    parts.extend(f'<th style="padding:7px 10px;text-align:left">{html.escape(str(c))}</th>' for c in frame.columns)
    parts.append('</tr></thead><tbody>')
    for i, row in enumerate(frame.to_dict("records")):
        parts.append('<tr>')
        for col, value in row.items():
            style = "padding:7px 10px;"
            if col == "player" and "signal" in frame and i < len(teams) and frame["player"].eq(frame.iloc[i]["player"]).sum() > 1:
                bg, fg = TEAMS.get(teams[i], (PAPER, "111111"))
                style += f"background:#{bg};color:#{fg};"
            if col == "signal":
                style += "font-weight:bold;"
            text = "" if grouped_name(frame, i, col) else display_value(value, col)
            parts.append(f'<td style="{style}">{html.escape(text)}</td>')
        parts.append('</tr>')
    parts.append('</tbody></table></div>')
    return "".join(parts)


def png_bytes(frame: pd.DataFrame) -> bytes:
    from PIL import Image, ImageDraw, ImageFont
    regular = "/System/Library/Fonts/Menlo.ttc"
    font = ImageFont.truetype(regular, 22)
    bold = ImageFont.truetype(regular, 22, index=1)
    rows = [[display_value(v, c) for c, v in row.items()] for row in frame.to_dict("records")]
    widths = [int(max([bold.getlength(str(col))] + [font.getlength(row[i]) for row in rows])) + 28
              for i, col in enumerate(frame.columns)]
    height = 40 * (len(rows) + 1) + 16
    image = Image.new("RGB", (sum(widths) + 16, height), "#" + PAPER)
    draw = ImageDraw.Draw(image)
    x = 8
    for col, width in zip(frame.columns, widths):
        draw.text((x + 10, 8), str(col), font=bold, fill="#111111")
        x += width
    teams = frame.attrs.get("teams", [])
    for i, values in enumerate(rows):
        x, y = 8, 40 * (i + 1)
        for j, (col, width) in enumerate(zip(frame.columns, widths)):
            fg = "111111"
            if col == "player" and "signal" in frame and i < len(teams) and frame["player"].eq(frame.iloc[i]["player"]).sum() > 1:
                bg, fg = TEAMS.get(teams[i], (PAPER, "111111"))
                draw.rectangle((x, y, x + width, y + 40), fill="#" + bg)
            value = "" if grouped_name(frame, i, col) else values[j]
            draw.text((x + 10, y + 7), value, font=bold if col == "signal" else font, fill="#" + fg)
            x += width
    out = BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def excel_bytes(frame: pd.DataFrame, research: pd.DataFrame | None = None) -> bytes:
    research = frame if research is None else research
    try:
        import openpyxl
    except ImportError:
        runtime = configured_path(
            "MODEL4_EXCEL_PYTHON",
            Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3",
        )
        if not runtime.is_file() or runtime.resolve() == Path(sys.executable).resolve():
            raise RuntimeError(
                "Excel export requires openpyxl. Install requirements.txt in the active Python "
                "environment, or set MODEL4_EXCEL_PYTHON to an interpreter with openpyxl."
            ) from None
        payload = json.dumps({"columns": list(frame.columns), "rows": frame.fillna("").values.tolist(),
                              "teams": frame.attrs.get("teams", []), "research_columns": list(research.columns),
                              "research_rows": research.fillna("").values.tolist()}, default=str).encode()
        result = subprocess.run([str(runtime), "-B", "-m", __name__], input=payload, capture_output=True,
                                cwd=Path(__file__).resolve().parents[2], timeout=60)
        if result.returncode:
            raise RuntimeError("Excel export failed: " + result.stderr.decode()[-1000:])
        return result.stdout
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Posting View"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(list(frame.columns))
    teams = frame.attrs.get("teams", [])
    for i, record in enumerate(frame.to_dict("records")):
        values = ["" if grouped_name(frame, i, col) else value for col, value in record.items()]
        sheet.append(values)
    for row in sheet:
        for cell in row:
            col = str(frame.columns[cell.column - 1])
            cell.font = Font(name="Menlo", size=11, bold=cell.row == 1 or col == "signal", color="111111")
            cell.fill = PatternFill("solid", fgColor=PAPER)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if cell.row > 1:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                value = number(cell.value)
                if value is not None and col not in {"player", "signal", "signal_history", "player_name", "opposing_pitcher"}:
                    cell.value = value
                    cell.number_format = '+0;-0;0' if col in {"book_odds", "model_odds", "parlay_odds"} else (
                        '0.00' if decimal_places(col) == 2 else '0' if value.is_integer() else '0.0')
                if col == "player" and "signal" in frame and cell.row - 2 < len(teams) and frame["player"].eq(frame.iloc[cell.row - 2]["player"]).sum() > 1:
                    bg, fg = TEAMS.get(teams[cell.row - 2], (PAPER, "111111"))
                    cell.fill = PatternFill("solid", fgColor=bg)
                    cell.font = Font(name="Menlo", size=11, color=fg)
        sheet.row_dimensions[row[0].row].height = 23
    for i, col in enumerate(frame.columns, 1):
        length = max([len(str(col))] + [len(display_value(v, col)) for v in frame[col]])
        sheet.column_dimensions[get_column_letter(i)].width = length * 1.22 + 3
    raw = book.create_sheet("Research Data")
    raw.append(list(research.columns))
    for row in research.fillna("").values.tolist():
        raw.append(row)
    for row in raw:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    raw.freeze_panes = "A2"
    raw.auto_filter.ref = raw.dimensions
    out = BytesIO()
    book.save(out)
    return out.getvalue()


if __name__ == "__main__":
    payload = json.load(sys.stdin)
    frame = pd.DataFrame(payload["rows"], columns=payload["columns"])
    frame.attrs["teams"] = payload.get("teams", [])
    research = pd.DataFrame(payload.get("research_rows", payload["rows"]), columns=payload.get("research_columns", payload["columns"]))
    sys.stdout.buffer.write(excel_bytes(frame, research))
